"""Resume parsing: extract raw text from PDF, DOCX, and TXT resumes."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path


@dataclass
class ParsedResume:
    source_file: str
    raw_text: str
    candidate_name: str | None = None
    email: str | None = None
    phone: str | None = None


EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(\+?\d[\d\s().-]{8,}\d)")


def parse_resume(path: str | Path) -> ParsedResume:
    """Parse a resume file (.pdf, .docx, .txt) into text + light metadata."""
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        text = _extract_pdf(path)
    elif suffix == ".docx":
        text = _extract_docx(path)
    elif suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type: {suffix} (supported: .pdf, .docx, .txt)")

    text = _normalize(text)
    return ParsedResume(
        source_file=str(path),
        raw_text=text,
        candidate_name=_guess_name(text),
        email=_first(EMAIL_RE, text),
        phone=_first(PHONE_RE, text),
    )


def parse_directory(directory: str | Path) -> list[ParsedResume]:
    """Parse every supported resume in a directory."""
    directory = Path(directory)
    resumes = []
    for path in sorted(directory.iterdir()):
        if path.suffix.lower() in {".pdf", ".docx", ".txt", ".md"}:
            resumes.append(parse_resume(path))
    return resumes


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _normalize(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _guess_name(text: str) -> str | None:
    """Heuristic: the first short, title-cased line is usually the name."""
    for line in text.splitlines()[:5]:
        line = line.strip()
        if 1 < len(line.split()) <= 4 and len(line) < 50 and not EMAIL_RE.search(line):
            if all(w[0].isupper() for w in line.split() if w[0].isalpha()):
                return line
    return None


def _first(pattern: re.Pattern, text: str) -> str | None:
    match = pattern.search(text)
    return match.group(0).strip() if match else None
